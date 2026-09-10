#!/usr/bin/env python3
"""Generate the editable GraphX architecture DOCX from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GraphX_Architecture.md"
OUTPUT = ROOT / "docs" / "GraphX_Architecture.docx"
ASSETS = ROOT / "docs" / "architecture-assets"
VERSION = (ROOT / "VERSION").read_text(encoding="ascii").strip()

NAVY = "17324D"
BLUE = "2563A6"
TEAL = "16877A"
LIGHT_BLUE = "EAF2FA"
LIGHT_TEAL = "E8F5F2"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "687482"
WHITE = "FFFFFF"
BLACK = "111820"


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    if bold:
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ]
    for path in candidates:
        if path and Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, subtitle="", fill="#EAF2FA", outline="#2563A6", width=3):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = xy
    draw.text(((x1 + x2) / 2, y1 + 23), title, anchor="mm", fill="#17324D",
              font=font(28, True))
    if subtitle:
        lines = subtitle.split("\n")
        y = y1 + 55
        for line in lines:
            draw.text(((x1 + x2) / 2, y), line, anchor="mm", fill="#465564", font=font(20))
            y += 25


def arrow(draw, start, end, color="#60788F", width=5):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    for delta in (2.55, -2.55):
        point = (end[0] + length * math.cos(angle + delta),
                 end[1] + length * math.sin(angle + delta))
        draw.line([end, point], fill=color, width=width)


def canvas(title: str):
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 55), title, fill="#17324D", font=font(42, True))
    draw.line((70, 112, 1730, 112), fill="#16877A", width=5)
    return image, draw


def create_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)

    image, d = canvas("GraphX architectural planes")
    rounded_box(d, (70, 165, 410, 295), "Authoritative model", "graphx.yaml + strict validation", "#E8F5F2", "#16877A")
    plane_titles = ["Logical graph", "Transport", "Network", "Deployment", "Observability", "Control / GUI"]
    # Paint connectors before their destination boxes so long diagonal routes
    # cannot obscure labels inside the architectural-plane nodes.
    for i, title in enumerate(plane_titles):
        x = 500 + (i % 3) * 410
        y = 155 + (i // 3) * 175
        arrow(d, (410, 230), (x, y + 60))
    for i, title in enumerate(plane_titles):
        x = 500 + (i % 3) * 410
        y = 155 + (i // 3) * 175
        rounded_box(d, (x, y, x + 330, y + 120), title, "peer configuration view")
    rounded_box(d, (110, 565, 470, 740), "Runtime data plane", "processes • containers • external nodes\nGraphX frames or raw protocols", "#F3F5F7", "#687482")
    rounded_box(d, (570, 565, 930, 740), "Infrastructure plane", "bridge • macvlan • ipvlan • OVS\nrouter • policy • mirror • netem", "#F3F5F7", "#687482")
    rounded_box(d, (1030, 565, 1390, 740), "Observation plane", "live state • metrics • OTLP\nhistory • PCAPNG", "#F3F5F7", "#687482")
    rounded_box(d, (1450, 565, 1730, 740), "Presentation", "GUI • Grafana\nWireshark", "#F3F5F7", "#687482")
    arrow(d, (470, 650), (570, 650)); arrow(d, (930, 650), (1030, 650)); arrow(d, (1390, 650), (1450, 650))
    d.text((70, 880), "Key rule: correlation joins the planes; lifecycle and trust boundaries remain explicit.", fill="#465564", font=font(28, True))
    image.save(ASSETS / "architecture-planes.png", quality=95)

    image, d = canvas("Mixed macvlan / IPvlan reference path")
    boxes = [
        ((50, 315, 270, 485), "Generator", "10.10.0.10\nexplicit MAC", "#E8F5F2", "#16877A"),
        ((320, 315, 540, 485), "macvlan L2", "gx-mac-domain", "#EAF2FA", "#2563A6"),
        ((590, 315, 810, 485), "OVS", "br-gx-mac\nSPAN: cap-mac", "#EAF2FA", "#2563A6"),
        ((860, 265, 1110, 535), "Router netns", "10.10.0.1 / 10.20.0.1\nforwarding + nftables\nnetem fault hook", "#FFF4DC", "#B87900"),
        ((1160, 315, 1380, 485), "OVS", "br-gx-ipv\nSPAN: cap-ipv", "#EAF2FA", "#2563A6"),
        ((1430, 315, 1650, 485), "IPvlan L2", "gx-ipv-domain", "#EAF2FA", "#2563A6"),
    ]
    for box, title, subtitle, fill, outline in boxes:
        rounded_box(d, box, title, subtitle, fill, outline)
    for a, b in zip(boxes, boxes[1:]):
        arrow(d, (a[0][2], (a[0][1]+a[0][3])//2), (b[0][0], (b[0][1]+b[0][3])//2))
    rounded_box(d, (1380, 675, 1580, 825), "Transform", "10.20.0.20", "#E8F5F2", "#16877A")
    rounded_box(d, (1600, 675, 1770, 825), "Sink", "10.20.0.30", "#E8F5F2", "#16877A")
    arrow(d, (1540, 485), (1480, 675)); arrow(d, (1580, 750), (1600, 750))
    d.text((65, 650), "samples edge: crosses two L2 domains and a router", fill="#17324D", font=font(27, True))
    d.text((65, 705), "transformed edge: stays inside the IPvlan domain", fill="#17324D", font=font(27, True))
    d.text((65, 875), "Native Linux is the semantic reference. Docker Desktop substitutes bridge domains and userspace OVS.", fill="#465564", font=font(24))
    image.save(ASSETS / "mixed-network-path.png", quality=95)

    image, d = canvas("IPvlan L2: three switched domains and one router")
    rounded_box(d, (45, 310, 245, 470), "Generator", "10.41.1.10", "#E8F5F2", "#16877A")
    rounded_box(d, (285, 310, 500, 470), "Generator L2", "IPvlan + OVS\nSPAN", "#EAF2FA", "#2563A6")
    rounded_box(d, (655, 230, 985, 550), "Namespace router", "10.41.1.1 / 10.41.2.1 / 10.41.3.1\nforwarding + nftables policies", "#FFF4DC", "#B87900")
    rounded_box(d, (1140, 155, 1365, 315), "Transform L2", "IPvlan + OVS\nSPAN", "#EAF2FA", "#2563A6")
    rounded_box(d, (1510, 155, 1735, 315), "Transform", "10.41.2.20", "#E8F5F2", "#16877A")
    rounded_box(d, (1140, 555, 1365, 715), "Sink L2", "IPvlan + OVS\nSPAN", "#EAF2FA", "#2563A6")
    rounded_box(d, (1510, 555, 1735, 715), "Sink", "10.41.3.30", "#E8F5F2", "#16877A")
    arrow(d, (245, 390), (285, 390)); arrow(d, (500, 390), (655, 390))
    arrow(d, (985, 330), (1140, 235)); arrow(d, (1365, 235), (1510, 235))
    arrow(d, (985, 450), (1140, 635)); arrow(d, (1365, 635), (1510, 635))
    d.text((70, 820), "Logical samples path: generator domain → router → transform domain", fill="#17324D", font=font(25, True))
    d.text((70, 865), "Logical transformed path: transform domain → router → sink domain", fill="#17324D", font=font(25, True))
    image.save(ASSETS / "ipvlan-l2-path.png", quality=95)

    image, d = canvas("One QEMU application, three deployment profiles")
    profiles = (
        (55, "External compatibility", "Host QEMU + slirp", "macOS or Linux\nQMP + guest probes"),
        (630, "Container compatibility", "QEMU container", "non-root; KVM or TCG\nslirp + relay"),
        (1205, "M6 TAP and OVS", "Non-root QEMU", "owned TAP + VLAN\nQMP + OVS SPAN"),
    )
    for offset, heading, middle_title, middle_sub in profiles:
        d.text((offset, 165), heading, fill="#17324D", font=font(27, True))
        rounded_box(d, (offset, 230, offset+190, 380), "Peer", "TCP + UDP", "#E8F5F2", "#16877A")
        rounded_box(d, (offset+225, 210, offset+520, 400), middle_title, middle_sub, "#EAF2FA", "#2563A6")
        arrow(d, (offset+190, 305), (offset+225, 305))
        rounded_box(d, (offset+55, 515, offset+465, 695), "Passive packet observer", "network_packet telemetry\nEthernet PCAPNG + packet SQLite", "#F3F5F7", "#687482")
        arrow(d, (offset+370, 400), (offset+300, 515))
    d.text((75, 830), "Same x86_64 guest and TCP/UDP application • M6 moves the primary data plane to GraphX-owned OVS", fill="#465564", font=font(27, True))
    image.save(ASSETS / "qemu-profiles.png", quality=95)

    image, d = canvas("Observation evidence to operator interfaces")
    rounded_box(d, (55, 260, 340, 430), "Evidence sources", "runtime callbacks\nQEMU packet observer\nOVS SPAN", "#E8F5F2", "#16877A")
    rounded_box(d, (430, 245, 750, 445), "Telemetry service", "normalize + bound\nlive state + SLO\nauthorization", "#EAF2FA", "#2563A6")
    arrow(d, (340, 345), (430, 345))
    destinations = [
        ((850, 145, 1170, 285), "GraphX GUI", "Application • Network\nHistory • Capture"),
        ((850, 330, 1170, 470), "Prometheus / Grafana", "metrics • dashboards • alerts"),
        ((850, 515, 1170, 655), "SQLite history", "bounded metadata\nseparate packet history"),
        ((850, 700, 1170, 840), "OTLP receiver", "bounded HTTPS export"),
    ]
    for box, title, subtitle in destinations:
        rounded_box(d, box, title, subtitle, "#F3F5F7", "#687482")
        arrow(d, (750, 345), (box[0], (box[1]+box[3])//2))
    rounded_box(d, (1290, 245, 1720, 445), "PCAPNG evidence", "USER0 GraphX frames\nEthernet packets", "#FFF4DC", "#B87900")
    rounded_box(d, (1290, 565, 1720, 765), "Analysis and download", "GUI catalog/download\nWireshark Lua • extcap • TShark", "#FFF4DC", "#B87900")
    arrow(d, (1505, 445), (1505, 565))
    d.text((75, 885), "Live telemetry is metadata and best effort. Raw captures retain payload bytes and require separate protection.", fill="#465564", font=font(26, True))
    image.save(ASSETS / "observability-gui-flow.png", quality=95)

    image, d = canvas("Architecture to verification evidence")
    stages = [
        ((45, 305, 305, 485), "Architecture + ADR", "boundaries\nand decisions", "#E8F5F2", "#16877A"),
        ((350, 305, 610, 485), "Example model", "graphx.yaml\nand assets", "#EAF2FA", "#2563A6"),
        ((655, 305, 915, 485), "Demo lifecycle", "start • verify\nobserve • stop", "#FFF4DC", "#B87900"),
        ((960, 305, 1220, 485), "Test layers", "unit • portable\nnative • quality", "#F3F5F7", "#687482"),
        ((1265, 305, 1755, 485), "Verification report", "commit • host • logs\nevidence • skips • cleanup", "#E8F5F2", "#16877A"),
    ]
    for box, title, subtitle, fill, outline in stages:
        rounded_box(d, box, title, subtitle, fill, outline)
    for first, second in zip(stages, stages[1:]):
        arrow(d, (first[0][2], 395), (second[0][0], 395))
    rounded_box(d, (330, 650, 1470, 815), "Shared implementation under test",
                "configuration • runtime • infrastructure • telemetry • control • history • capture • GUI",
                "#F3F5F7", "#687482")
    arrow(d, (785, 485), (785, 650))
    arrow(d, (1090, 650), (1090, 485))
    d.text((70, 890), "Profiles aggregate evidence; they do not change what a platform or test can prove.",
           fill="#465564", font=font(27, True))
    image.save(ASSETS / "architecture-evidence-flow.png", quality=95)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part: continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); set_run_font(run, "Aptos Mono", 8.5, color="153B57")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(part); set_run_font(run)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(document, rows):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    available = 6.65
    widths = [available / len(rows[0])] * len(rows[0])
    if len(rows[0]) == 2: widths = [2.0, 4.65]
    elif len(rows[0]) == 3: widths = [1.8, 2.65, 2.2]
    elif len(rows[0]) == 4: widths = [1.35, 2.05, 1.55, 1.7]
    for ridx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if ridx == 0: set_repeat_table_header(row)
        for cidx, value in enumerate(values):
            cell = row.cells[cidx]
            cell.width = Inches(widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx == 0: shade(cell, NAVY)
            elif ridx % 2 == 0: shade(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            for run in p.runs:
                set_run_font(run, size=8.0 if len(rows[0]) >= 4 else 8.5,
                             bold=(ridx == 0), color=WHITE if ridx == 0 else BLACK)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(document, code):
    p = document.add_paragraph(style="Code")
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    shade_paragraph = OxmlElement("w:shd"); shade_paragraph.set(qn("w:fill"), LIGHT_GRAY)
    p._p.get_or_add_pPr().append(shade_paragraph)
    run = p.add_run(code.rstrip())
    set_run_font(run, "Aptos Mono", 8.2, color="17324D")


def add_figure(document, path, caption):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.7))
    cap = document.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.72); section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82); section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3); section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(9.6)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color, before, after in (
        ("Title", 31, BLACK, 0, 14), ("Heading 1", 19, BLACK, 15, 7),
        ("Heading 2", 14, BLACK, 11, 5), ("Heading 3", 11.5, BLACK, 9, 4)):
        style = styles[name]
        style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    # Let Word flow major sections naturally. A forced page break on every
    # section creates nearly empty pages when a preceding section ends late.
    styles["Heading 1"].paragraph_format.page_break_before = False
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Caption"].font.name = "Aptos"; styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.italic = True; styles["Caption"].font.color.rgb = RGBColor.from_string(MID_GRAY)
    if "Code" not in styles:
        styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)

    header = section.header.paragraphs[0]
    header.text = "GRAPHX  |  ARCHITECTURE AND NETWORK TOPOLOGY"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs: set_run_font(run, size=7.5, bold=True, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])
    for run in section.footer.paragraphs[0].runs: set_run_font(run, size=8, color=MID_GRAY)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    review_date = next(
        line.removeprefix("**Review date:** ")
        for line in lines
        if line.startswith("**Review date:** ")
    )
    # Cover page
    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(80); p.add_run(title)
    sub = doc.add_paragraph("A living reference for logical graphs, transports, network infrastructure, QEMU integration, observability, capture, history, control, and GUI behavior")
    sub.paragraph_format.space_before = Pt(12); sub.paragraph_format.space_after = Pt(26)
    for run in sub.runs: set_run_font(run, size=15, color=BLACK)
    band = doc.add_table(rows=1, cols=1); band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0); shade(cell, NAVY); set_cell_margins(cell, 220, 220, 220, 220)
    cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(cp, f"GraphX {VERSION}\nRepository architecture baseline\nReviewed {review_date}")
    for run in cp.runs: set_run_font(run, size=12, bold=True, color=WHITE)
    doc.add_paragraph()
    note = doc.add_paragraph("Document status: maintained source and editable Word edition. Implemented behavior is separated from proposed work.")
    for run in note.runs: set_run_font(run, size=10, color=MID_GRAY)
    doc.add_page_break()

    # Compact contents (major sections only)
    doc.add_heading("Contents", level=1)
    for line in lines:
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line[3:])
    doc.add_page_break()

    # Cover already carries the source metadata; resume at the first section.
    i = next(index for index, line in enumerate(lines) if line == "## Executive summary")
    mermaid_index = 0
    figure_map = [
        (ASSETS / "architecture-planes.png", "Figure 1. GraphX configuration views and runtime planes."),
        (ASSETS / "ipvlan-l2-path.png", "Figure 2. IPvlan L2 domains connected by OVS and a namespace router."),
        (ASSETS / "mixed-network-path.png", "Figure 3. Mixed-network logical path and explicit observation/fault points."),
        (ASSETS / "observability-gui-flow.png", "Figure 5. Live, durable, exported, and captured evidence presented to operators."),
        (ASSETS / "architecture-evidence-flow.png", "Figure 6. Traceability from architecture through examples, demos, tests, and verification."),
    ]
    while i < len(lines):
        raw = lines[i]; stripped = raw.strip()
        if not stripped:
            i += 1; continue
        if stripped.startswith("```mermaid"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"): i += 1
            path, caption = figure_map[mermaid_index]
            add_figure(doc, path, caption); mermaid_index += 1; i += 1; continue
        if stripped.startswith("```"):
            i += 1; code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            add_code(doc, "\n".join(code)); i += 1; continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i); add_table(doc, rows); continue
        if raw.startswith("## "):
            doc.add_heading(raw[3:], level=1)
            i += 1; continue
        if raw.startswith("### "):
            heading = raw[4:]
            doc.add_heading(heading, level=2)
            if heading.startswith("8.1 "):
                add_figure(doc, ASSETS / "qemu-profiles.png", "Figure 4. Shared QEMU application with three deployment ownership profiles.")
            i += 1; continue
        if raw.startswith("#### "):
            doc.add_heading(raw[5:], level=3); i += 1; continue
        if re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            add_inline(p, stripped); i += 1; continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, stripped[2:]); i += 1; continue
        if stripped.startswith("**") and stripped.endswith("**") and ":" in stripped:
            p = doc.add_paragraph(); add_inline(p, stripped); i += 1; continue
        # Join ordinary wrapped Markdown lines into one paragraph.
        paragraph = [stripped]; i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", "```", "- ")) or re.match(r"^\d+\. ", nxt)):
                break
            paragraph.append(nxt); i += 1
        p = doc.add_paragraph(); add_inline(p, " ".join(paragraph))

    props = doc.core_properties
    props.title = "GraphX Architecture and Network Topology"
    props.subject = (
        f"GraphX {VERSION} system, network, QEMU, observability, capture, history, "
        "control, and GUI architecture"
    )
    props.author = "GraphX Project"
    props.keywords = "GraphX, architecture, network topology, Docker, Open vSwitch, QEMU, PCAPNG, observability"
    doc.save(OUTPUT)


if __name__ == "__main__":
    create_diagrams()
    build_docx()
    print(OUTPUT)
